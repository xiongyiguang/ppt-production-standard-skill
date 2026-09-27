"""Regression for retired entry and open arrow endpoints."""
import unittest
from pptx import Presentation
from build_layout_library import Canvas,build
class RetiredLibraryTests(unittest.TestCase):
    def test_retired_entry(self):
        with self.assertRaises(RuntimeError):build('institute','unused.pptx')
    def test_open_arrows_in_all_directions(self):
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);c=Canvas(s)
        ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        for coords in [(0,0,10,0),(10,0,0,0),(0,0,0,10),(0,10,0,0)]:
            end=c.edge(*coords)._element.find('.//a:tailEnd',ns)
            self.assertEqual(end.get('type'),'arrow')
            self.assertEqual((end.get('w'),end.get('len')),('med','med'))
        self.assertIsNone(c.edge(0,0,10,0,False)._element.find('.//a:tailEnd',ns))
if __name__=='__main__':unittest.main(verbosity=2)
