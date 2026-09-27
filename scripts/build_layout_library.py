"""Native drawing primitives retained for existing project reproduction; old layout library generation is retired.

Metadata is attributed to Wise PPT in references/layout-library/NOTICE.md.
The drawing code uses PowerPoint shapes/charts and preserves the template package.
"""
from __future__ import annotations
import argparse, copy, hashlib, json, math, zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
BRANDS = {
    'richinfo': ('彩讯股份2026版-20260725-封面页、内容页、过渡页、结尾页模板.pptx', '2_内容1'),
    'mobile': ('新建PPT模板-中国移动主题色规范修正版.pptx', '内容页1'),
    'institute': ('新建PPT模板-中国移动主题色规范修正版 - 创新研究院.pptx', '内容页1'),
    'legacy': ('新建PPT模板-彩讯科技主题色规范修正版V2.pptx', '内容页1'),
}
COLORS = {'ink': MSO_THEME_COLOR.DARK_1, 'paper': MSO_THEME_COLOR.LIGHT_1,
          'muted': MSO_THEME_COLOR.DARK_2, 'light': MSO_THEME_COLOR.LIGHT_2,
          'accent': MSO_THEME_COLOR.ACCENT_1, 'contrast': MSO_THEME_COLOR.ACCENT_4}

def color(target, token):
    target.theme_color = COLORS[token]

def text_format(shape, text, size=14, bold=False, ink='ink', align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Cm(.16)
    tf.margin_top = tf.margin_bottom = Cm(.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    # Conservative width estimate; actual PowerPoint line counts are verified separately.
    width_pt = shape.width / 12700 - 10
    multiline = '\n' in text or sum(1 if ord(c) > 255 else .55 for c in text) * size > width_pt
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_before = Pt(0); p.space_after = Pt(0)
        p.line_spacing = 1.2 if multiline else 1.0
        r = p.add_run(); r.text = line; r.font.name = 'Arial'; r.font.size = Pt(size)
        r.font.bold = bold; color(r.font.color, ink)
        rp = r._r.get_or_add_rPr()
        ea = OxmlElement('a:ea'); ea.set('typeface', 'Microsoft YaHei'); rp.append(ea)
    return shape

class Canvas:
    def __init__(self, slide):
        self.s = slide
    def xy(self,x,y,w=0,h=0):
        return [Cm(v) for v in (1.35+x*.3115, 4.05+y*.123, w*.3115,h*.123)]
    def box(self,x,y,w,h,text='',fill='light',size=14,bold=False,kind=MSO_SHAPE.RECTANGLE,ink='ink'):
        sh=self.s.shapes.add_shape(kind,*self.xy(x,y,w,h))
        sh._element.spPr.append(OxmlElement('a:effectLst'))
        sh.name='library: '+(text.replace('\n',' / ')[:50] or 'geometry')
        if fill is None: sh.fill.background()
        else: sh.fill.solid(); color(sh.fill.fore_color,fill)
        sh.line.fill.background()
        if text: text_format(sh,text,size,bold,ink)
        return sh
    def label(self,x,y,w,h,text,size=14,bold=False,ink='ink',center=False):
        sh=self.s.shapes.add_textbox(*self.xy(x,y,w,h)); sh.name='library: '+text[:50]
        return text_format(sh,text,size,bold,ink,PP_ALIGN.CENTER if center else PP_ALIGN.LEFT)
    def edge(self,x1,y1,x2,y2,arrow=True,token='muted'):
        x,y,_,_=self.xy(x1,y1);xx,yy,_,_=self.xy(x2,y2)
        sh=self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x,y,xx,yy)
        sh._element.spPr.append(OxmlElement('a:effectLst'))
        sh.name='library: relationship';color(sh.line.color,token);sh.line.width=Pt(1.5)
        if arrow:
            end=OxmlElement('a:tailEnd');end.set('type','arrow');end.set('w','med');end.set('len','med');sh._element.spPr.get_or_add_ln().append(end)
        return sh
    def card(self,x,y,w,h,title,body='',accent=False):
        sh=self.box(x,y,w,h,title+('\n'+body if body else ''),'accent' if accent else 'light',14,False,ink='paper' if accent else 'ink')
        # Keep titles inside their containers; never overlay a second text box.
        first=sh.text_frame.paragraphs[0]
        for run in first.runs:run.font.bold=True;run.font.size=Pt(16)
        return sh
    def media(self,x,y,w,h,label='真实素材'):
        sh=self.box(x,y,w,h,label+'\n请插入授权原件','light')
        sh.name='asset-slot: '+label;return sh
    def row(self,labels,y=35,x=0,w=100,h=26,arrows=True):
        n=len(labels);gap=3;cw=(w-gap*(n-1))/n
        for i,t in enumerate(labels):
            self.card(x+i*(cw+gap),y,cw,h,t,accent=i==n-1)
            if arrows and i<n-1:self.edge(x+i*(cw+gap)+cw,y+h/2,x+(i+1)*(cw+gap),y+h/2)
    def grid(self,labels,cols=3,x=0,y=0,w=100,h=92):
        rows=math.ceil(len(labels)/cols);gx=2;gy=4;cw=(w-gx*(cols-1))/cols;ch=(h-gy*(rows-1))/rows
        for i,t in enumerate(labels):self.card(x+(i%cols)*(cw+gx),y+(i//cols)*(ch+gy),cw,ch,t,accent=i==0)
    def chart(self,x,y,w,h,kind='bar',names=None,values=None):
        names=names or ['样本甲','样本乙','样本丙'];values=values or [20,35,45]
        d=CategoryChartData();d.categories=names;d.add_series('演示数据',values)
        typ={'bar':XL_CHART_TYPE.COLUMN_CLUSTERED,'horizontal':XL_CHART_TYPE.BAR_CLUSTERED,'line':XL_CHART_TYPE.LINE,'pie':XL_CHART_TYPE.DOUGHNUT}[kind]
        ch=self.s.shapes.add_chart(typ,*self.xy(x,y,w,h),d).chart
        ch.has_title=False;ch.has_legend=kind=='pie'
        ch.font.name='Arial';ch.font.size=Pt(12);color(ch.font.color,'ink')
        for ser in ch.series:
            ser.format.fill.solid();color(ser.format.fill.fore_color,'accent')
            color(ser.format.line.color,'accent')
            if kind=='pie':
                for i,p in enumerate(ser.points):p.format.fill.solid();color(p.format.fill.fore_color,['accent','contrast','light'][i%3])
        if kind=='pie':ch.legend.position=XL_LEGEND_POSITION.BOTTOM
        else:
            ch.category_axis.tick_labels.font.name='Microsoft YaHei';ch.category_axis.tick_labels.font.size=Pt(12)
            ch.value_axis.tick_labels.font.name='Arial';ch.value_axis.tick_labels.font.size=Pt(12)
            ch.value_axis.has_major_gridlines=True
            color(ch.value_axis.major_gridlines.format.line.color,'light')
            ch.value_axis.major_gridlines.format.line.width=Pt(.5)
        return ch

def protected(name):
    return name.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/','ppt/notesMasters/','ppt/handoutMasters/'))

def build(*args,**kwargs):
    raise RuntimeError('The 75-layout library is retired. Use content-led design; do not generate old layouts.')

if __name__=='__main__':
    raise SystemExit('Retired: no layout-library generation entry is available.')
